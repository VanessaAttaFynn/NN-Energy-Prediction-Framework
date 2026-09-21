"""WP-A: rebalance the research-question hierarchy in the thesis .docx.

Applies the approved change set:

  1. Global remap of internal cross-references (Section / RQ / Table) to the new
     numbering, done simultaneously so no mapping collides.
  2. Reorder the Results (S5) and Discussion (S6) sections so the strong
     within-family result leads and the transferability probe comes last.
  3. Remove the rejected "Exploratory Probe" labels from headings and RQs.
  4. Rewrite the abstract, the contributions list, the research questions, and
     the opening paragraphs of the merged transferability sections.
  5. Restore the deleted "1.3 Contributions" heading.
  6. Insert a new S6.1 discussing the primary within-family result, and a
     paragraph on the physical mechanism behind cross-family failure.
  7. Strip the leftover `font-claude-response-body` style from S1.1.

Reads and writes distinct paths; the input file is never modified.

Usage:
    python scripts/apply_wpa_edits.py INPUT.docx OUTPUT.docx
"""
from __future__ import annotations

import copy
import re
import sys

from docx import Document
from docx.text.paragraph import Paragraph

# --------------------------------------------------------------------------
# New numbering maps. Each is applied in a single regex pass so that, e.g.,
# "Section 5.3" -> "Section 5.5" and "Section 5.5" -> "Section 5.3" cannot
# interfere with one another.
# --------------------------------------------------------------------------
SECTION_MAP = {
    "Sections 5.3\u20135.4": "Section 5.5",
    # Identity mapping: this range means "the discussion of all four research
    # questions", which is still 6.1-6.4 after the reorder. Listed explicitly
    # because the longer key takes precedence over the bare "Sections 6.1" rule
    # below, which would otherwise rewrite it to "Sections 6.4 through 6.4".
    "Sections 6.1 through 6.4": "Sections 6.1 through 6.4",
    "Sections 6.1": "Sections 6.4",
    "Section 5.3": "Section 5.5",
    "Section 5.4": "Section 5.5",
    "Section 5.5": "Section 5.3",
    "Section 5.6": "Section 5.4",
    "Section 6.1": "Section 6.4",
    "Section 6.2": "Section 6.4",
    "Section 6.3": "Section 6.2",
    "Section 6.4": "Section 6.3",
}

RQ_MAP = {"RQ1": "RQ4", "RQ2": "RQ4", "RQ3": "RQ2", "RQ4": "RQ3"}

TABLE_MAP = {
    "Table 7b": "Table 8",
    "Table 10": "Table 9",
    "Table 11": "Table 10",
    "Table 12": "Table 11",
    "Table 8": "Table 12",
    "Table 9": "Table 13",
}


def remap(text: str, mapping: dict[str, str]) -> str:
    """Apply every mapping key in one simultaneous pass."""
    pattern = re.compile("|".join(re.escape(k) for k in sorted(mapping, key=len, reverse=True)))
    return pattern.sub(lambda m: mapping[m.group(0)], text)


# --------------------------------------------------------------------------
# Replacement prose. Written in the NEW numbering, so it is inserted after the
# global remap has already run.
# --------------------------------------------------------------------------
ABSTRACT_BACKGROUND = (
    "Background: Neural architecture search and other large-scale training campaigns consume "
    "substantial energy, yet no a priori method for estimating training energy before a run "
    "begins has been validated beyond the single architecture family and single hardware "
    "platform on which it was developed. A systematic review of six existing a priori methods "
    "found that none has been tested for generalisation across architecture families or "
    "hardware platforms, and none addresses the measurement-instrument disagreement documented "
    "between hardware watt-meter and software-based estimation tools. This work investigates how "
    "accurately a family-agnostic feature representation can predict training energy within an "
    "architecture family, and probes what would be required to extend that prediction across "
    "architecture families and hardware platforms, using two independently collected "
    "training-energy datasets covering convolutional and multilayer perceptron architectures."
)

ABSTRACT_METHODS = (
    "Methods: A Random Forest regression model was trained on a feature set restricted to "
    "properties measurable identically across architecture families, supplemented by "
    "family-specific auxiliary features validated through ablation. Two datasets, BUTTER-E "
    "(multilayer perceptrons, hardware-measured) and EC-NAS (convolutional networks, "
    "software-measured), were combined under three training conditions to separate "
    "architecture-family effects from measurement-instrument effects. The model was evaluated "
    "under a standard random split, first within each architecture family and then on the pooled "
    "data, with its generalisation to unseen training datasets validated by leave-one-dataset-out "
    "cross-validation. A literature-derived correction for measurement bias and a "
    "power-times-duration output decomposition were each tested against their alternatives. As a "
    "secondary investigation, the same model was then tested zero-shot on an architecture family "
    "and on hardware withheld entirely from training."
)

ABSTRACT_RESULTS = (
    "Results: Within each architecture family, the model achieved a Kendall-Tau rank correlation "
    "of 0.861 to 0.936, exceeding both a FLOPs-only baseline and a re-implemented prior surrogate "
    "model on identical data, and this held up under leave-one-dataset-out cross-validation across "
    "all twelve of BUTTER-E's source datasets. Dataset-level properties \u2014 chiefly the number "
    "of training observations \u2014 rather than architectural properties, were found to drive a "
    "substantial share of multilayer-perceptron training-energy variance. Measurement-bias "
    "correction produced no change in aggregate fit metrics while shifting absolute predicted "
    "energy by approximately twenty-five percent, and decomposing energy into power and duration "
    "gave no measurable accuracy advantage over predicting it directly. Testing the same model "
    "zero-shot on an architecture family and on hardware withheld entirely from training, a "
    "condition no prior a priori method has reported, failed for specific and independently "
    "diagnosed reasons, which are carried into the future-work directions in Section 7."
)

CONTRIBUTIONS_LEAD_IN = (
    "This paper's primary contribution is an a priori training-energy predictor validated within "
    "each of two architecture families, together with a finding about what actually drives that "
    "energy. A secondary strand then tests a harder question no prior method has attempted, and "
    "reports what was found:"
)

CONTRIBUTION_BULLETS = [
    "A training-energy predictor for multilayer perceptron architectures (via BUTTER-E) and "
    "convolutional architectures (via EC-NAS), each validated against re-implemented prior-style "
    "baselines on identical data and substantially outperforming them within its own family "
    "(Section 5.1).",

    "A finding, validated by leave-one-dataset-out cross-validation, that the scale of the dataset "
    "being trained on is a dominant and largely unreported driver of multilayer-perceptron training "
    "energy, distinct from the architectural properties every reviewed a priori method treats as "
    "the sole predictors (Section 5.2).",

    "An empirical test of two design questions with no precedent in the reviewed literature: "
    "whether a literature-derived measurement-bias correction affects cross-dataset pooling (RQ2), "
    "and whether decomposing energy into power and duration outperforms predicting it directly "
    "(RQ3).",

    "A first, zero-shot test of transfer to an entirely unseen architecture family and hardware "
    "platform, with mechanistic diagnosis of why it fails under the data currently available "
    "(RQ4).",
]

RESEARCH_QUESTIONS = [
    "RQ1 (within-family prediction): How accurately can training energy consumption be predicted "
    "before a run begins, using only architectural, configuration and dataset information available "
    "at that point, and how does that accuracy compare with FLOPs-based and prior surrogate-model "
    "baselines?",

    "RQ2 (measurement-bias correction): Does correcting software-measured power values for known "
    "instrument underestimation improve joint-model accuracy relative to no correction?",

    "RQ3 (output structure): Does predicting power and duration separately (E = P \u00d7 D) "
    "outperform predicting energy directly as a single scalar target, given identical features and "
    "training data?",

    "RQ4 (transferability): Does a model trained within one architecture family and on one hardware "
    "platform retain meaningful accuracy on an architecture family or hardware platform withheld "
    "entirely from training, and if not, why?",
]

TRANSFERABILITY_OPENING = (
    "RQ4 asks whether a model trained within one architecture family and on one hardware platform "
    "retains meaningful accuracy on an architecture family or hardware platform withheld entirely "
    "from training. It is addressed in two parts, both using only the five core, family-agnostic "
    "features specified in Section 4.1: because the auxiliary features validated in Section 5.2 are "
    "by construction defined for one architecture family each, using them here would test the model "
    "on features it could not, even in principle, have learned to use. The architecture direction "
    "trains on each family in turn and tests on the whole of the other (Table 12). The hardware "
    "direction uses EC-NAS's hardware-specific benchmark, in which ninety-one architectures were "
    "independently measured on four distinct graphics processing units, training exclusively on the "
    "Quadro RTX 6000 measurements and testing on each of the remaining three GPU classes (Table 13)."
)

DISCUSSION_RQ1 = (
    "Within each architecture family the model predicts training energy with a Kendall-Tau rank "
    "correlation of 0.936 for BUTTER-E and 0.861 for EC-NAS, outperforming both baselines under an "
    "identical feature set (Section 5.1). Two features of this result deserve comment. First, the "
    "same five core features perform very differently against the weak baseline in the two families: "
    "FLOPs alone attains a coefficient of determination of 0.609 for EC-NAS but only 0.087 for "
    "BUTTER-E. That asymmetry is the clearest single indication that architectural scale is not a "
    "sufficient account of training energy in the multilayer perceptron case, and it is what "
    "motivated the feature investigation in Section 5.2. Second, the accuracy achieved here rests "
    "substantially on a feature every reviewed a priori method omits: the scale of the dataset being "
    "trained on. Once that is included, in a form that remains defined for datasets unseen during "
    "training, multilayer perceptron prediction reaches parity with the convolutional case. The "
    "within-family result is therefore best read not as evidence that the core architectural feature "
    "set is sufficient on its own, but that it becomes sufficient once the cost of the data being "
    "processed is accounted for alongside it."
)

DISCUSSION_TRANSFER_MERGED = (
    "The answer is no in both directions, and each is established with more precision than a simple "
    "negative result would ordinarily allow. In the architecture direction, Section 5.5 identifies "
    "two independent mechanisms, one operating in each direction of the test, rather than a single "
    "undifferentiated failure of transfer, though this design cannot fully isolate architecture "
    "family from the other factors named in Section 6.5 that vary alongside it. When EC-NAS is used "
    "as the sole training family, the model never observes variation in the auxiliary features shown "
    "in Section 5.2 to carry substantial predictive information for BUTTER-E, since those features "
    "are masked to zero throughout EC-NAS's training data; the resulting failure is a coverage gap "
    "rather than a failure of the model to learn a transferable relationship. When BUTTER-E is used "
    "as the sole training family, a fixed training-epoch budget of three thousand, present throughout "
    "BUTTER-E's measured data, leaves the model unable to represent how energy scales with training "
    "length, and this single mismatch alone is sufficient to produce the severe absolute-scale "
    "miscalibration observed in that direction, despite the model's rank ordering of architectures "
    "remaining comparatively well preserved. In the hardware direction, the same conclusion is "
    "reached through two independent approaches rather than one. A learned model trained on a single "
    "GPU class exhibits no evidence of genuine hardware-conditioned behaviour: its predictions for "
    "unseen GPUs track the training GPU's own energy scale regardless of the target hardware's actual "
    "characteristics, and the one target GPU for which this model performs well, Titan Xp, does so "
    "because that GPU's true energy profile for these architectures happens to closely resemble the "
    "training GPU's, not because the model has learned any functional relationship between hardware "
    "specification and energy consumption. Supplying the model with explicit hardware-specification "
    "features did not correct this behaviour, for the same structural reason identified above: those "
    "features are constant throughout a single-GPU training set and carry no variance for the model "
    "to learn from."
)

PHYSICAL_MECHANISM = (
    "Underlying both diagnosed mechanisms is a more general observation about what the core feature "
    "set actually measures. Parameter count, depth and FLOPs describe the computational work an "
    "architecture performs; they do not describe how the hardware performs it. Energy consumption in "
    "a training run is the product of a power draw and a duration, and the power draw at any moment "
    "depends on which units are active, how well the workload occupies them, and how much of the time "
    "is spent on memory movement rather than arithmetic. Two architectures with identical parameter "
    "counts can differ substantially in all three, and the difference is larger, not smaller, when "
    "the two architectures perform structurally different kinds of operation. The failure of "
    "zero-shot transfer across families is therefore not best understood as a failure of the model, "
    "but as a mismatch between the quantity the features describe and the quantity energy actually "
    "depends on. This is consistent with the within-family results, where the feature that closes "
    "most of the residual gap \u2014 the scale of the dataset being trained on \u2014 is likewise a "
    "proxy for work performed, but one that happens to align closely with how duration, and therefore "
    "energy, scales within a family."
)

CONCLUSION_TRANSFER = (
    "Beyond this primary result, a secondary investigation asked whether the same approach could "
    "generalise beyond the single architecture family and single hardware platform within which every "
    "method identified in the preceding systematic literature review was developed and validated "
    "\u2014 a zero-shot test no prior a priori method has attempted or reported. It failed in both "
    "directions, and both failures are supported by specific, independently diagnosed mechanisms "
    "rather than left as unexplained shortfalls. Section 6.4 identifies feature-coverage gaps, "
    "training-configuration range mismatches, and the absence of hardware variation in available "
    "training data as the proximate causes, and each is a property of the data available to this work "
    "rather than an inherent limitation of the modelling approach taken. This diagnostic work is "
    "offered as a contribution in its own right, since it converts an unattempted question into a "
    "specific, actionable account of what conditions a future dataset or evaluation design would need "
    "to satisfy before this question could be answered differently \u2014 not as a second primary "
    "result standing alongside Section 5.1's positive finding."
)

# --------------------------------------------------------------------------
# Heading texts, keyed by the paragraph index in the INPUT document.
# --------------------------------------------------------------------------
S51_HEADING = "5.1 RQ1: Within-Family A Priori Prediction Performance"
S52_HEADING = "5.2 Feature Validation: The Role of Dataset Scale and Hardware Type"
S53_HEADING = "5.3 RQ2: Measurement-Bias Correction"          # was 5.5
S54_HEADING = "5.4 RQ3: Output Structure Comparison"           # was 5.6
S55_HEADING = "5.5 RQ4: Architecture and Hardware Transferability"  # merged 5.3 + 5.4
S61_HEADING = "6.1 RQ1: How accurately is training energy predicted within an architecture family?"
S62_HEADING = "6.2 RQ2: Does correcting for measurement-instrument bias improve pooled model performance?"  # was 6.3
S63_HEADING = "6.3 RQ3: Does decomposing energy into power and duration improve prediction?"               # was 6.4
S64_HEADING = "6.4 RQ4: Does the model transfer to an unseen architecture family or hardware platform?"    # merged 6.1 + 6.2

# Input paragraph indices (verified against scripts/inspect_docx.py).
I_ABSTRACT_BG, I_ABSTRACT_METH, I_ABSTRACT_RES = 8, 9, 10
I_PROBLEM_HEADING = 21          # "1.2 Problem statement" - heading template
I_CONTRIB_LEAD_IN = 25
I_CONTRIB_BULLETS = [26, 27, 28, 29]
I_RQ = [31, 32, 33, 34]
I_S51, I_S52 = 130, 138
I_S53, I_S54, I_S55, I_S56 = 151, 166, 177, 181
I_S6, I_S7 = 188, 208
I_S61, I_S62, I_S63, I_S64, I_S65 = 189, 193, 197, 200, 202
I_CONCLUSION_TRANSFER = 212


def set_text(par: Paragraph, text: str) -> None:
    """Replace a paragraph's text, preserving the first run's formatting."""
    if not par.runs:
        par.add_run(text)
        return
    par.runs[0].text = text
    for run in par.runs[1:]:
        run.text = ""


def iter_all_paragraphs(doc):
    """Body paragraphs plus every paragraph inside every table."""
    for par in doc.paragraphs:
        yield par
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    yield par


def clone_before(anchor: Paragraph, template: Paragraph, text: str) -> Paragraph:
    """Insert a copy of `template` immediately before `anchor`, with new text."""
    new_el = copy.deepcopy(template._p)
    anchor._p.addprevious(new_el)
    new_par = Paragraph(new_el, anchor._parent)
    set_text(new_par, text)
    return new_par


def block_before(body, start: Paragraph, stop: Paragraph) -> list:
    """Body-level elements from `start` up to (not including) `stop`."""
    children = list(body)
    return children[children.index(start._p): children.index(stop._p)]


def move_block_before(body, elements: list, anchor: Paragraph) -> None:
    for el in elements:
        el.getparent().remove(el)
    for el in elements:
        anchor._p.addprevious(el)


def main() -> None:
    src, dst = sys.argv[1], sys.argv[2]
    doc = Document(src)
    body = doc.element.body
    paras = doc.paragraphs
    log: list[str] = []

    # -- capture object references before anything moves --------------------
    abstract = [paras[i] for i in (I_ABSTRACT_BG, I_ABSTRACT_METH, I_ABSTRACT_RES)]
    problem_heading = paras[I_PROBLEM_HEADING]
    contrib_lead = paras[I_CONTRIB_LEAD_IN]
    contrib_bullets = [paras[i] for i in I_CONTRIB_BULLETS]
    rqs = [paras[i] for i in I_RQ]
    s51, s52 = paras[I_S51], paras[I_S52]
    s53, s54, s55, s56 = paras[I_S53], paras[I_S54], paras[I_S55], paras[I_S56]
    s6, s7 = paras[I_S6], paras[I_S7]
    s61, s62, s63, s64, s65 = (paras[i] for i in (I_S61, I_S62, I_S63, I_S64, I_S65))
    conclusion_transfer = paras[I_CONCLUSION_TRANSFER]

    assert s51.text.startswith("5.1"), s51.text
    assert s53.text.startswith("5.3"), s53.text
    assert s61.text.startswith("6.1"), s61.text
    assert s6.text.startswith("6."), s6.text
    assert s7.text.startswith("7."), s7.text

    # -- 1. global cross-reference remap ------------------------------------
    for par in iter_all_paragraphs(doc):
        original = par.text
        if not original:
            continue
        updated = remap(remap(remap(original, SECTION_MAP), RQ_MAP), TABLE_MAP)
        if updated != original:
            set_text(par, updated)
            log.append(f"remap: {original[:60]!r} -> {updated[:60]!r}")

    # -- 2. reorder the Results section -------------------------------------
    # current: 5.1 5.2 [5.3+5.4] [5.5 5.6]
    # target:  5.1 5.2 [5.5 5.6] [5.3+5.4]
    bias_and_output = block_before(body, s55, s6)
    move_block_before(body, bias_and_output, s53)
    log.append(f"moved Results block ({len(bias_and_output)} elements) before the transferability section")

    # -- 3. reorder the Discussion section ----------------------------------
    # current: [6.1+6.2] [6.3 6.4] 6.5
    # target:  [6.3 6.4] [6.1+6.2] 6.5
    bias_and_output_disc = block_before(body, s63, s65)
    move_block_before(body, bias_and_output_disc, s61)
    log.append(f"moved Discussion block ({len(bias_and_output_disc)} elements) before 6.1")

    # -- 4. heading and prose replacements ----------------------------------
    for par, text, label in (
        (abstract[0], ABSTRACT_BACKGROUND, "abstract/background"),
        (abstract[1], ABSTRACT_METHODS, "abstract/methods"),
        (abstract[2], ABSTRACT_RESULTS, "abstract/results"),
        (contrib_lead, CONTRIBUTIONS_LEAD_IN, "1.3/lead-in"),
        (s51, S51_HEADING, "heading 5.1"),
        (s52, S52_HEADING, "heading 5.2"),
        (s53, S55_HEADING, "heading 5.5"),
        (s55, S53_HEADING, "heading 5.3"),
        (s56, S54_HEADING, "heading 5.4"),
        (s61, S64_HEADING, "heading 6.4"),
        (s63, S62_HEADING, "heading 6.2"),
        (s64, S63_HEADING, "heading 6.3"),
        (conclusion_transfer, CONCLUSION_TRANSFER, "7.1/transfer paragraph"),
    ):
        set_text(par, text)
        log.append(f"set {label}")

    for par, text in zip(contrib_bullets, CONTRIBUTION_BULLETS):
        set_text(par, text)
    log.append("set 4 contribution bullets")

    for par, text in zip(rqs, RESEARCH_QUESTIONS):
        set_text(par, text)
    log.append("set 4 research questions")

    # 5.5 opening: replace the old architecture-direction opening, drop the
    # old hardware-direction heading and opening (now folded into 5.5).
    transfer_body = paras[152]
    set_text(transfer_body, TRANSFERABILITY_OPENING)
    log.append("set 5.5 opening")

    for heading in (s54,):
        heading._p.getparent().remove(heading._p)
    log.append("removed the old 5.4 heading")

    hw_opening = paras[167]
    hw_opening._p.getparent().remove(hw_opening._p)
    log.append("removed the old hardware opening (folded into 5.5)")

    # 6.4: merge the two old openings into the first, drop the second heading
    # and the duplicate opening paragraph.
    arch_opening = paras[190]
    set_text(arch_opening, DISCUSSION_TRANSFER_MERGED)
    log.append("merged the 6.4 opening")

    s62._p.getparent().remove(s62._p)
    log.append("removed the old 6.2 heading")

    paras[194]._p.getparent().remove(paras[194]._p)
    log.append("removed the duplicate hardware opening paragraph")

    # -- 5. restore the deleted 1.3 heading ---------------------------------
    clone_before(contrib_lead, problem_heading, "1.3 Contributions")
    log.append("restored the 1.3 Contributions heading")

    # -- 6. new 6.1, and the physical-mechanism paragraph -------------------
    # Inserted before s63 (the old measurement-bias heading), which is the
    # first section of the reordered Discussion, so that 6.1 leads.
    body_par_template = paras[191]          # an ordinary body paragraph
    heading_template = s63                  # a Heading 2
    new_61_heading = clone_before(s63, heading_template, S61_HEADING)
    clone_before(s63, body_par_template, DISCUSSION_RQ1)
    log.append("inserted new 6.1 (heading + discussion of the primary result)")

    clone_before(s65, body_par_template, PHYSICAL_MECHANISM)
    log.append("inserted the physical-mechanism paragraph at the end of 6.4")

    # -- 7. fix a pre-existing broken cross-reference -----------------------
    # "discussed further in Section 6.3" in the transferability discussion
    # pointed at the measurement-bias section; the conditions named are
    # discussed in Future Work.
    fixed = 0
    for par in doc.paragraphs:
        if "discussed further in Section 6.2" in par.text:
            set_text(par, par.text.replace("discussed further in Section 6.2",
                                           "discussed further in Section 7.3"))
            fixed += 1
    log.append(f"repointed {fixed} broken cross-reference (Section 6.2 -> 7.3)")

    # -- 8. strip the leftover AI-tool paragraph style ----------------------
    stripped = 0
    for par in doc.paragraphs:
        if par.style.name == "font-claude-response-body":
            par.style = doc.styles["Normal"]
            stripped += 1
    log.append(f"reset {stripped} paragraphs from font-claude-response-body to Normal")

    doc.save(dst)
    print("\n".join(log))
    print(f"\nsaved: {dst}")
    print(f"paragraphs now: {len(doc.paragraphs)}, tables: {len(doc.tables)}, "
          f"images: {len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
