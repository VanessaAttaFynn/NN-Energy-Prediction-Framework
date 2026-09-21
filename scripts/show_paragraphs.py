"""Print the full text of a range of paragraphs from a .docx.

Usage:
    python scripts/show_paragraphs.py FILE START END
"""
import sys

from docx import Document


def main() -> None:
    doc = Document(sys.argv[1])
    start, end = int(sys.argv[2]), int(sys.argv[3])
    for i in range(start, min(end, len(doc.paragraphs))):
        par = doc.paragraphs[i]
        text = par.text.strip()
        if text:
            print(f"--- [{i}] ({par.style.name}) ---")
            print(text)
            print()


if __name__ == "__main__":
    main()
