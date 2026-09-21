"""Dump the structure of a .docx: paragraph index, style, and text.

Usage:
    python scripts/inspect_docx.py "docs/write-ups/Thesis_Assembled (1).docx" [--all]

Without --all, prints only headings and short numbered/uppercase lines plus the
first 40 paragraphs. With --all, prints every paragraph.
"""
import sys

from docx import Document


def main() -> None:
    path = sys.argv[1]
    show_all = "--all" in sys.argv

    doc = Document(path)
    print(f"### file: {path}")
    print(f"### paragraphs: {len(doc.paragraphs)}")
    print(f"### tables: {len(doc.tables)}")
    print(f"### sections: {len(doc.sections)}")
    print(f"### inline shapes (images): {len(doc.inline_shapes)}")
    print()

    for i, par in enumerate(doc.paragraphs):
        text = par.text.strip()
        style = par.style.name

        is_heading = style.lower().startswith("heading") or style.lower() == "title"
        is_short_numbered = text and len(text) < 90 and text[0].isdigit()
        is_upper = text and len(text) < 90 and text.isupper()

        if show_all or i < 40 or is_heading or is_short_numbered or is_upper:
            print(f"{i:5d} [{style:16s}] {text[:130]}")


if __name__ == "__main__":
    main()
